from docx import Document
import openpyxl
import re

# 📄 Word-Dokument laden
doc = Document("TEST20250617 Entwurf Arbeitshilfe Prüfung DORA clean.docx")

# ✅ Erwartete Spaltenüberschriften
erwartete_header = ["Themenkomplex", "Beispielhafte Prüfungshandlungen"]

# 🔍 Funktion: alle Tabellen mit passenden Headern finden
def finde_alle_passenden_tabellen(doc, erwartete_header):
    passende_tabellen = []
    for table in doc.tables:
        if len(table.rows) == 0:
            continue
        erste_zeile = [cell.text.strip() for cell in table.rows[0].cells]
        if all(any(h.lower() in z.lower() for z in erste_zeile) for h in erwartete_header):
            passende_tabellen.append(table)
    return passende_tabellen

# 🧠 DORA-Artikel extrahieren
def extrahiere_dora_artikel(text):
    artikel = re.findall(r"Artikel\s+\d+(?:\s+Abs\.\s*\d+)?\s+DORA", text)
    return sorted(set(artikel))

# 🔧 Prüfungshandlungen in zwei Kategorien aufteilen
def teile_pruefungshandlungen(text):
    angemessenheit = []
    wirksamkeit = []
    aktueller_block = None

    for zeile in text.splitlines():
        zeile = zeile.strip()
        if zeile.lower().startswith("prüfung der angemessenheit"):
            aktueller_block = "angemessenheit"
        elif zeile.lower().startswith("prüfung der wirksamkeit"):
            aktueller_block = "wirksamkeit"
        elif zeile.startswith("•") or zeile.startswith("-") or zeile.startswith("–"):
            punkt = re.sub(r"^[•\-\–]\s*", "", zeile).strip()
            if aktueller_block == "angemessenheit":
                angemessenheit.append(punkt)
            elif aktueller_block == "wirksamkeit":
                wirksamkeit.append(punkt)
    return angemessenheit, wirksamkeit

# 📊 Excel-Datei vorbereiten
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "DORA-Prüfungen"
ws.append(["Themenkomplex", "Angemessenheit", "Wirksamkeit", "DORA-Artikel"])

# 📥 Tabellen extrahieren und verarbeiten
tabellen = finde_alle_passenden_tabellen(doc, erwartete_header)
if not tabellen:
    print("❌ Keine passenden Tabellen gefunden.")
    exit()

for tabelle in tabellen:
    for row in tabelle.rows[1:]:  # Erste Zeile = Header
        zellen = [cell.text.strip() for cell in row.cells]
        if len(zellen) < 2:
            continue

        themenkomplex = zellen[0]
        pruefungshandlungen = zellen[1]

        angemessenheit_punkte, wirksamkeit_punkte = teile_pruefungshandlungen(pruefungshandlungen)
        artikel_text = "\n".join(angemessenheit_punkte + wirksamkeit_punkte)
        dora_artikel = ", ".join(extrahiere_dora_artikel(artikel_text))

        ws.append([
            themenkomplex,
            "\n".join(angemessenheit_punkte),
            "\n".join(wirksamkeit_punkte),
            dora_artikel
        ])

# ✨ Formatierung: Zeilenumbruch aktivieren
for row in ws.iter_rows(min_row=2, max_col=4):
    for cell in row:
        cell.alignment = openpyxl.styles.Alignment(wrap_text=True)

# 📐 Spaltenbreiten automatisch setzen
for col in ws.columns:
    max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
    col_letter = col[0].column_letter
    ws.column_dimensions[col_letter].width = min(max_len + 5, 80)

# 💾 Datei speichern
wb.save("DORA_Pruefungsauswertung.xlsx")
print("✅ Fertig! Datei gespeichert als 'DORA_Pruefungsauswertung.xlsx'")
