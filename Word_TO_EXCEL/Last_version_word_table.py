from docx import Document
import openpyxl
import re

# 📄 Word-Dokument laden
doc = Document("TEST20250617 Entwurf Arbeitshilfe Prüfung DORA clean.docx")

# ✅ Gewünschte Spaltennamen (sollen mindestens enthalten sein)
erforderliche_spalten = ["Themenkomplex", "Beispielhafte Prüfungshandlungen", "DORA"]

# 📊 Excel vorbereiten
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Alle Tabellen"

# 🔧 Funktion für vollständigen Textinhalt einer Zelle
def get_full_cell_text(cell):
    return "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip())

# 🔍 Prüfen, ob Tabelle alle benötigten Spalten enthält
def enthaelt_erforderliche_spalten(headerzeile, erforderlich):
    header_normiert = [h.strip().lower() for h in headerzeile]
    return all(any(erw.lower() in h for h in header_normiert) for erw in erforderlich)

# 📥 Tabellen durchlaufen und in Excel schreiben
zeilenzahl = 1

for table in doc.tables:
    if len(table.rows) < 2:
        continue  # zu kleine Tabelle überspringen

    # Header extrahieren
    headerzellen = [get_full_cell_text(cell) for cell in table.rows[0].cells]
    if not enthaelt_erforderliche_spalten(headerzellen, erforderliche_spalten):
        continue

    # Header in Excel schreiben
    ws.append(headerzellen)
    zeilenzahl += 1

    # Datenzeilen schreiben
    for row in table.rows[1:]:
        zellen = [get_full_cell_text(cell) for cell in row.cells]
        if all(z == "" for z in zellen):
            continue
        # Mit leeren Zellen auffüllen
        while len(zellen) < len(headerzellen):
            zellen.append("")
        ws.append(zellen)
        zeilenzahl += 1

    # Leerzeile zwischen Tabellen
    ws.append([])
    zeilenzahl += 1

# ✨ Zeilenumbruch aktivieren
for row in ws.iter_rows(min_row=1, max_row=zeilenzahl):
    for cell in row:
        cell.alignment = openpyxl.styles.Alignment(wrap_text=True)

# 📐 Spaltenbreite automatisch anpassen
for col in ws.columns:
    max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
    col_letter = col[0].column_letter
    ws.column_dimensions[col_letter].width = min(max_len + 5, 80)

# 💾 Excel-Datei speichern
wb.save("DORA_Alle_Tabellen_Vollständig.xlsx")
print("✅ Excel-Datei 'DORA_Alle_Tabellen_Vollständig.xlsx' wurde erfolgreich erstellt.")
